"""The ONS's Price Index of Private Rents (dataset `ons-private-rents`): the average monthly
private rent in pounds for the UK, its countries, regions and local authorities. Open Government
Licence v3.0.

The release is pinned: each month's file has its own address, and a build reads the same file
until the release is changed here.
"""

import datetime
from pathlib import Path

from openpyxl import load_workbook

from atlas.fetch import download

RELEASE = "21october2026"
NAME = "priceindexofprivaterentsukmonthlypricestatistics"
URL = (
    "https://www.ons.gov.uk/file?uri=/economy/inflationandpriceindices/datasets/"
    f"{NAME}/{RELEASE}/{NAME}.xlsx"
)
SHEET = "Table 1"


def fetch(data_dir: Path) -> Path:
    """Downloads the workbook to `data/raw/ons-pipr/` once (about 18 MB)."""
    target = data_dir / "raw" / "ons-pipr" / "pipr.xlsx"
    download(URL, target)
    return target


def read_rents(path: Path) -> dict[str, float]:
    """The average monthly rent (pounds) of each area, by ONS code, in its latest month that has
    a price."""
    book = load_workbook(path, read_only=True, data_only=True)
    latest: dict[str, tuple[datetime.datetime, float]] = {}
    header_seen = False
    for row in book[SHEET].iter_rows(values_only=True):
        if not header_seen:
            header_seen = row[0] == "Time period"
            continue
        month, code, price = row[0], row[1], row[7]
        if not isinstance(month, datetime.datetime) or not isinstance(price, int | float):
            continue
        if not isinstance(code, str):
            continue
        if code not in latest or month > latest[code][0]:
            latest[code] = (month, float(price))
    book.close()
    return {code: price for code, (_, price) in latest.items()}
