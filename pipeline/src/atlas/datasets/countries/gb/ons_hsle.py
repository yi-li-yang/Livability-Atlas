"""The ONS's health state life expectancy estimates for middle layer super output areas (dataset
`ons-health-state-life-expectancy`): years of life in good health at birth, by sex, for the MSOAs of
England and Wales, 2019 to 2023. Open Government Licence v3.0.

The estimates join deaths registered over five years to the 2021 census's answers on health. The
release is pinned: each release has its own address.
"""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from atlas.fetch import download

RELEASE = "2019to2023"
NAME = "healthstatelifeexpectancyestimatesformiddlelayersuperoutputareasenglandandwales"
URL = (
    "https://www.ons.gov.uk/file?uri=/peoplepopulationandcommunity/healthandsocialcare/"
    f"healthandlifeexpectancies/datasets/{NAME}/{RELEASE}/hslemsoa.xlsx"
)
SHEET = "2"  # healthy life expectancy at birth, by sex


def fetch(data_dir: Path) -> Path:
    """Downloads the workbook (about 3 MB) to `data/raw/ons-hsle/` once, and returns it."""
    target = data_dir / "raw" / "ons-hsle" / "hslemsoa.xlsx"
    download(URL, target)
    return target


def read_msoa(path: Path) -> pd.Series:
    """The healthy life expectancy at birth in years of each MSOA, indexed by its code: the mean
    of the estimates for men and for women."""
    book = load_workbook(path, read_only=True, data_only=True)
    values: dict[str, list[float]] = {}
    columns: dict[str, int] = {}
    for row in book[SHEET].iter_rows(values_only=True):
        if not columns:
            if row and row[0] == "Period":
                columns = {str(name): i for i, name in enumerate(row)}
            continue
        if row[columns["Area type"]] != "MSOA":
            continue
        value = row[columns["HLE"]]
        if isinstance(value, int | float):
            values.setdefault(str(row[columns["Area code"]]), []).append(float(value))
    return pd.Series({code: sum(v) / len(v) for code, v in values.items()}, dtype=float)
